from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"(?<!\*)\*([^*\n]+?)\*(?!\*)")
_MATRIX_INLINE_RE = re.compile(r"\[\[.*?\]\]")
_LATEX_REPLACEMENTS = {
    r"\rightarrow": "→",
    r"\to": "→",
    r"\leftarrow": "←",
    r"\leftrightarrow": "↔",
    r"\times": "×",
    r"\cdot": "·",
    r"\leq": "≤",
    r"\geq": "≥",
    r"\neq": "≠",
    r"\pm": "±",
    r"\infty": "∞",
    r"\ldots": "...",
}

# OMML (Word math) namespace
_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

# \mathbb{R} dan kawan-kawan → karakter matematika double-struck
_MATHBB_MAP = {
    "R": "ℝ", "Z": "ℤ", "N": "ℕ", "Q": "ℚ", "C": "ℂ", "H": "ℍ",
    "P": "ℙ", "E": "𝔼", "F": "𝔽", "A": "𝔸", "B": "𝔹", "D": "𝔻",
}


def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _omml_run(text: str) -> str:
    return f"<m:r><m:t>{_esc(text)}</m:t></m:r>"


def _omml_sup(base: str, exp: str) -> str:
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{exp}</m:sup></m:sSup>"


def _omml_sub(base: str, sub: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"


def _omml_sqrt(inner: str) -> str:
    return (
        '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr>'
        f"<m:deg/><m:e>{inner}</m:e></m:rad>"
    )


def _omml_matrix(cell_rows: list[list[str]]) -> str:
    body = "".join(
        "<m:mr>" + "".join(f"<m:e>{cell}</m:e>" for cell in row) + "</m:mr>"
        for row in cell_rows
    )
    matrix = f"<m:m>{body}</m:m>"
    return (
        '<m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr>'
        f"<m:e>{matrix}</m:e></m:d>"
    )


def _omml_fraction(num: str, den: str, *, bar: bool = True) -> str:
    num_xml = "".join(_parse_latex_math(num)) or _omml_run("")
    den_xml = "".join(_parse_latex_math(den)) or _omml_run("")
    kind = "bar" if bar else "noBar"
    return (
        f"<m:f><m:fPr><m:type m:val=\"{kind}\"/></m:fPr>"
        f"<m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>"
    )


def _omml_sqrt_n(inner: str, deg: str = "") -> str:
    if deg:
        return (
            '<m:rad><m:radPr><m:degHide m:val="0"/></m:radPr>'
            f'<m:deg>{deg}</m:deg><m:e>{inner}</m:e></m:rad>'
        )
    return _omml_sqrt(inner)


def _omml_math_accent(cmd: str, inner: str) -> str:
    _ACCENT_CHR = {
        "hat": "̂", "widehat": "̂", "bar": "̄", "overline": "̄",
        "underline": "̲", "vec": "⃗", "dot": "̇", "ddot": "̈",
        "tilde": "̃", "widetilde": "̃", "check": "̌",
        "acute": "́", "grave": "̀", "breve": "̆",
    }
    inner_xml = "".join(_parse_latex_math(inner)) or _omml_run("")
    chr_ = _ACCENT_CHR.get(cmd, "")
    return (
        f'<m:acc><m:accPr><m:chr m:val="{chr_}"/></m:accPr>'
        f"<m:e>{inner_xml}</m:e></m:acc>"
    )


def _omml_nary(chr_: str, sub: str, sup: str, inner: str) -> str:
    sub_xml = "".join(_parse_latex_math(sub)) or _omml_run("")
    sup_xml = "".join(_parse_latex_math(sup)) or _omml_run("")
    inner_xml = "".join(_parse_latex_math(inner)) or _omml_run("")
    return (
        "<m:nary><m:naryPr>"
        f'<m:chr m:val="{chr_}"/><m:limLoc m:val="undOvr"/>'
        f'<m:subHide m:val="{"0" if sub else "1"}"/>'
        f'<m:supHide m:val="{"0" if sup else "1"}"/>'
        "</m:naryPr>"
        f"<m:sub>{sub_xml}</m:sub><m:sup>{sup_xml}</m:sup>"
        f"<m:e>{inner_xml}</m:e></m:nary>"
    )


def _parse_latex_math(s: str) -> list[str]:
    """Parse LaTeX math string into OMML XML element list.

    Handles: \\begin{bmatrix}...\\end{bmatrix}, \\sqrt{}, x^{n},
    \\alpha, \\times, \\cdot, \\neq, \\leq, \\geq, \\text{}, and
    inline * as ×.  Unknown commands are rendered as plain text.
    """
    _ALPHA_MAP = {
        "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ",
        "epsilon": "ε", "theta": "θ", "lambda": "λ", "mu": "μ",
        "pi": "π", "sigma": "σ", "phi": "φ", "omega": "ω",
        "rho": "ρ", "tau": "τ", "psi": "ψ", "chi": "χ",
        "eta": "η", "kappa": "κ", "xi": "ξ", "zeta": "ζ",
        "Delta": "Δ", "Sigma": "Σ", "Omega": "Ω", "Pi": "Π",
        "Phi": "Φ", "Psi": "Ψ", "Theta": "Θ", "Lambda": "Λ",
    }
    _CMD_MAP = {
        "times": "×", "cdot": "·", "pm": "±", "mp": "∓",
        "leq": "≤", "geq": "≥", "neq": "≠", "approx": "≈",
        "equiv": "≡", "sim": "∼", "propto": "∝", "mid": "|",
        "le": "≤", "ge": "≥", "ne": "≠", "lt": "<", "gt": ">",
        "rightarrow": "→", "leftarrow": "←", "leftrightarrow": "↔",
        "Rightarrow": "⇒", "Leftarrow": "⇐", "Leftrightarrow": "⇔",
        "to": "→", "mapsto": "↦", "longrightarrow": "⟶",
        "longleftarrow": "⟵", "iff": "⟺", "%": "%",
        "infty": "∞", "partial": "∂", "nabla": "∇",
        "forall": "∀", "exists": "∃", "in": "∈", "notin": "∉",
        "subset": "⊂", "supset": "⊃", "subseteq": "⊆", "supseteq": "⊇",
        "cup": "∪", "cap": "∩",
        "emptyset": "∅", "ldots": "…", "cdots": "⋯", "vdots": "⋮",
        "ddots": "⋱", "hbar": "ℏ", "ell": "ℓ",
        "langle": "⟨", "rangle": "⟩", "surd": "√", "backslash": "∖",
        "circ": "∘", "bullet": "∙", "cdotp": "·", "colon": ":",
        "sum": "∑", "prod": "∏", "int": "∫", "iint": "∬", "iiint": "∭",
        "oint": "∮", "bigcup": "⋃", "bigcap": "⋂",
        "bigoplus": "⊕", "bigotimes": "⊗", "coprod": "∐",
        "lim": "lim", "limsup": "lim sup", "liminf": "lim inf",
        "left": "", "right": "", "newcommand": "", "operatorname": " ",
        "text": " ", "quad": " ", "qquad": "  ", "hline": "",
    }

    nodes: list[str] = []
    i, n = 0, len(s)

    def _read_braced(pos: int) -> tuple[str, int]:
        while pos < n and s[pos] == " ":
            pos += 1
        if pos < n and s[pos] == "{":
            depth, start = 1, pos + 1
            pos += 1
            while pos < n and depth:
                if s[pos] == "{":
                    depth += 1
                elif s[pos] == "}":
                    depth -= 1
                pos += 1
            return s[start : pos - 1], pos
        if pos < n:
            return s[pos], pos + 1
        return "", pos

    while i < n:
        c = s[i]

        if c == " " or c == "\n":
            i += 1
            continue

        # LaTeX command
        if c == "\\":
            i += 1
            if i >= n:
                nodes.append(_omml_run("\\"))
                break
            nxt = s[i]
            # Escaped special: \{ \} \_ \^ \* \% \$
            if nxt in ("{", "}", "_", "^", "*", "%", "$", "|"):
                nodes.append(_omml_run(nxt))
                i += 1
                continue
            # Spacing commands: \  \, \; \: \! 
            if nxt == " ":
                i += 1
                continue
            if nxt in ",;:!":
                nodes.append(_omml_run(" " if nxt != "!" else ""))
                i += 1
                continue
            # Command name
            if nxt.isalpha():
                j = i
                while j < n and s[j].isalpha():
                    j += 1
                cmd = s[i:j]
                i = j
                if cmd in _ALPHA_MAP:
                    nodes.append(_omml_run(_ALPHA_MAP[cmd]))
                elif cmd == "frac":
                    num, i = _read_braced(i)
                    den, i = _read_braced(i)
                    nodes.append(_omml_fraction(num, den))
                elif cmd == "binom":
                    num, i = _read_braced(i)
                    den, i = _read_braced(i)
                    inner = _omml_fraction(num, den, bar=False)
                    nodes.append(
                        '<m:d><m:dPr><m:begChr m:val="("/>'
                        '<m:endChr m:val=")"/></m:dPr>'
                        f"<m:e>{inner}</m:e></m:d>"
                    )
                elif cmd == "sqrt":
                    deg = None
                    if i < n and s[i] == "[":
                        end = s.find("]", i)
                        if end != -1:
                            deg = s[i + 1:end]
                            i = end + 1
                    inner, i = _read_braced(i)
                    inner_xml = "".join(_parse_latex_math(inner))
                    deg_xml = "".join(_parse_latex_math(deg)) if deg else ""
                    nodes.append(_omml_sqrt_n(inner_xml, deg_xml))
                elif cmd == "text":
                    inner, i = _read_braced(i)
                    nodes.append(_omml_run(inner))
                elif cmd == "operatorname":
                    inner, i = _read_braced(i)
                    nodes.append(_omml_run(inner))
                elif cmd == "mathbb":
                    inner, i = _read_braced(i)
                    ch = inner.strip()
                    nodes.append(_omml_run(_MATHBB_MAP.get(ch, ch)))
                elif cmd in ("mathrm", "mathbf", "mathit", "mathcal",
                             "mathtt", "mathsf", "mathscr", "mathnormal"):
                    inner, i = _read_braced(i)
                    nodes.append(_omml_run(inner))
                elif cmd in ("hat", "widehat", "bar", "overline", "underline",
                             "vec", "dot", "ddot", "tilde", "widetilde",
                             "check", "acute", "grave", "breve"):
                    inner, i = _read_braced(i)
                    nodes.append(_omml_math_accent(cmd, inner))
                elif cmd == "begin":
                    env, i = _read_braced(i)
                    content, i = _read_until_end(s, i, env)
                    if env.endswith("matrix"):
                        nodes.append(_parse_latex_matrix(content, env))
                    else:
                        nodes.append(_omml_run(content))
                elif cmd in _CMD_MAP:
                    rep = _CMD_MAP[cmd]
                    if rep:
                        nodes.append(_omml_run(rep))
                else:
                    nodes.append(_omml_run(cmd))
                continue
            # \ followed by non-alpha: treat as text
            nodes.append(_omml_run(nxt))
            i += 1
            continue

        # Superscript
        if c == "^":
            base = nodes.pop() if nodes else _omml_run("")
            i += 1
            exp, i = _read_braced(i)
            nodes.append(_omml_sup(base, "".join(_parse_latex_math(exp)) if exp else _omml_run("")))
            continue

        # Subscript
        if c == "_":
            base = nodes.pop() if nodes else _omml_run("")
            i += 1
            sub, i = _read_braced(i)
            sub_nodes = "".join(_parse_latex_math(sub)) if sub else _omml_run("")
            nodes.append(_omml_sub(base, sub_nodes))
            continue

        # Brace group
        if c == "{":
            depth, start = 1, i + 1
            i += 1
            while i < n and depth:
                if s[i] == "{":
                    depth += 1
                elif s[i] == "}":
                    depth -= 1
                i += 1
            inner = s[start : i - 1]
            nodes.extend(_parse_latex_math(inner))
            continue

        # Parenthesized group → <m:d> delimiter (so (AB)_{11} stays one base
        # AND inner LaTeX like \beta_0 is parsed recursively, not raw text).
        if c == "(":
            depth, j = 1, i + 1
            while j < n and depth:
                if s[j] == "(":
                    depth += 1
                elif s[j] == ")":
                    depth -= 1
                j += 1
            if depth == 0:
                inner = "".join(_parse_latex_math(s[i + 1 : j - 1])) or _omml_run("")
                nodes.append(
                    '<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/>'
                    f"</m:dPr><m:e>{inner}</m:e></m:d>"
                )
                i = j
            else:
                nodes.append(_omml_run(s[i:]))
                i = n
            continue

        # ASCII matrix: [[a, b], [c, d]]
        if s.startswith("[[", i):
            j = s.find("]]", i)
            if j != -1:
                inner = s[i + 2 : j]
                rows: list[list[str]] = []
                for part in re.split(r"\]\s*,\s*\[", inner):
                    cells = [c.strip() for c in part.split(",") if c.strip()]
                    rows.append(
                        ["".join(_parse_latex_math(c)) if c else _omml_run("") for c in cells]
                    )
                nodes.append(_omml_matrix(rows))
                i = j + 2
                continue
            nodes.append(_omml_run(s[i]))
            i += 1
            continue

        # Numeric literal
        m = re.match(r"\d+(?:\.\d+)?", s[i:])
        if m:
            nodes.append(_omml_run(m.group()))
            i += m.end()
            continue

        # Alphabetic variable (multi-letter run)
        if c.isalpha():
            j = i
            while j < n and s[j].isalpha():
                j += 1
            nodes.append(_omml_run(s[i:j]))
            i = j
            continue

        # * outside \-command → ×
        if c == "*":
            nodes.append(_omml_run("×"))
            i += 1
            continue

        # Everything else: parentheses, operators, brackets, etc.
        nodes.append(_omml_run(c))
        i += 1

    return nodes


def _read_until_end(s: str, i: int, env: str) -> tuple[str, int]:
    """Read content until \\end{env}, returning (content, new_index)."""
    tag = f"\\end{{{env}}}"
    j = s.find(tag, i)
    if j == -1:
        return s[i:], len(s)
    return s[i:j], j + len(tag)


def _parse_latex_matrix(content: str, env: str) -> str:
    """Convert LaTeX matrix content to OMML with bracket delimiters."""
    rows = re.split(r"\\\\|(?<!\\)\\(?!\\)", content)
    cell_rows: list[list[str]] = []
    for row in rows:
        cells = [c.strip() for c in row.split("&")]
        cell_rows.append(
            ["".join(_parse_latex_math(c)) if c else _omml_run("") for c in cells]
        )
    return _omml_matrix(cell_rows)


def _equation_omml(math_str: str) -> str:
    content = "".join(_parse_latex_math(math_str)) or _omml_run("")
    return f'<m:oMath xmlns:m="{_MATH_NS}">{content}</m:oMath>'


def _append_equation(paragraph, math_str: str):
    try:
        from lxml import etree

        el = etree.fromstring(_equation_omml(math_str).encode("utf-8"))
        paragraph._p.append(el)
    except Exception:  # noqa: BLE001
        paragraph.add_run(math_str)


def _add_runs(paragraph, text: str):
    """Terapkan **bold** dan *italic* sederhana ke satu paragraf."""
    text = _normalize_text(text)
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            _add_italic_runs(paragraph, text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.bold = True
        pos = m.end()
    _add_italic_runs(paragraph, text[pos:])


def _add_italic_runs(paragraph, text: str):
    """Terapkan *italic* dan konversi matriks inline [[...]] jadi equation."""
    text = _normalize_text(text)
    pos = 0
    for m in _MATRIX_INLINE_RE.finditer(text):
        if m.start() > pos:
            _add_italic_text(paragraph, text[pos:m.start()])
        _append_equation(paragraph, m.group(0))
        pos = m.end()
    _add_italic_text(paragraph, text[pos:])


def _add_italic_text(paragraph, text: str):
    """Terapkan *italic* dan konversi ruas LaTeX inline (\\beta, \\sigma^2, ...)
    yang ditulis tanpa $...$ menjadi equation."""
    pos = 0
    for start, end in _find_math_spans(text):
        if start > pos:
            _add_italic_plain(paragraph, text[pos:start])
        _append_equation(paragraph, text[start:end])
        pos = end
    _add_italic_plain(paragraph, text[pos:])


def _add_italic_plain(paragraph, text: str):
    pos = 0
    for m in _ITALIC_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        run = paragraph.add_run(m.group(1))
        run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _normalize_text(text: str) -> str:
    """Hilangkan artefak LaTeX/Markdown yang tidak boleh tampil mentah di Word."""
    for source, replacement in _LATEX_REPLACEMENTS.items():
        text = text.replace(source, replacement)
    text = re.sub(r"\\(?:left|right)\b", "", text)
    text = re.sub(r"\\text\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"(?<![A-Za-z])rightarrow(?![A-Za-z])", "→", text)
    text = re.sub(r"(?<![A-Za-z])leftarrow(?![A-Za-z])", "←", text)
    return text


_LATEX_CMD_RE = re.compile(r"\\[A-Za-z]+")

# Kata yang sah di dalam ekspresi matematika (nama fungsi, kata LaTeX).
_MATH_WORDS = {
    "alpha", "beta", "gamma", "delta", "epsilon", "zeta", "eta", "theta",
    "iota", "kappa", "lambda", "mu", "nu", "xi", "omicron", "pi", "rho",
    "sigma", "tau", "upsilon", "phi", "chi", "psi", "omega",
    "Normal", "Bernoulli", "Binomial", "Poisson", "Exponential", "Uniform",
    "Gaussian", "sim", "mid", "given", "s.t", "of", "and", "or",
    "sin", "cos", "tan", "cot", "sec", "csc", "arcsin", "arccos", "arctan",
    "sinh", "cosh", "tanh", "log", "ln", "exp", "lim", "sup", "inf", "max",
    "min", "Pr", "P", "E", "Var", "Cov", "Corr", "det", "dim", "mod", "Ker",
    "Im", "Re", "adj", "sum", "prod", "int", "frac", "sqrt", "times", "cdot",
    "left", "right", "quad", "text", "begin", "end", "pmatrix", "bmatrix",
    "gather", "eqnarray", "matrix", "mid", "vert", "Vert",
}


def _tokenize_rough(line: str) -> list[str]:
    return re.findall(r"[A-Za-z]+|\d+(?:\.\d+)?|[^\sA-Za-z0-9]+", line)


def _line_is_display_math(line: str) -> bool:
    """True bila satu baris (tanpa $) sebagian besar adalah ekspresi LaTeX,
    mis. `YmidX;∼;Normal(\\beta_0 + \\beta_1 x_1,\\; \\sigma^2)`."""
    line = (line or "").strip()
    if not line or "$" in line:
        return False  # sudah ditangani jalur $...$ / $$...$$
    if not _LATEX_CMD_RE.search(line):
        return False  # tidak ada perintah LaTeX → teks biasa
    toks = _tokenize_rough(line)
    if not toks:
        return False
    prose = sum(
        1 for t in toks
        if t.isalpha() and len(t) >= 4 and t not in _MATH_WORDS
    )
    # Maksimal ~20% token adalah kata prosa murni: kalimat seperti
    # `Jadi nilai yang dicari adalah \beta = 5` jangan dijadikan equation.
    return prose * 5 <= len(toks)


def _scan_math_run_end(text: str, j: int) -> int:
    """Perluas dari posisi j sampai batas akhir ruas matematika yang masuk akal."""
    n = len(text)
    depth = 0
    while j < n:
        c = text[j]
        if c == " " and depth == 0:
            k = j
            while k < n and text[k] == " ":
                k += 1
            m = re.match(r"[\\A-Za-z0-9_.^]+", text[k:k + 10])
            nxt = m.group(0) if m else ""
            if nxt and "\\" not in nxt and re.fullmatch(r"[A-Za-z]+", nxt):
                if nxt.lower() not in _MATH_WORDS:
                    return j
                j = k
                continue
            j = k
            continue
        if c in "([{":
            depth += 1
        elif c in ")]}":
            depth -= 1
            if depth < 0:
                return j
        j += 1
    return n


def _find_math_spans(text: str) -> list[tuple[int, int]]:
    """Temukan ruas inline LaTeX tanpa $ di dalam teks → list (start, end).
    Ruas diperluas ke kurung pembuka/penutup yang mengelilinginya, jadi
    `Normal(\\beta_0 + \\beta_1, \\sigma^2)` ikut ditangkap mencakup parens."""
    spans: list[tuple[int, int]] = []
    i = 0
    n = len(text)
    while i < n:
        m = _LATEX_CMD_RE.search(text, i)
        if not m:
            break
        start = m.start()
        end = _scan_math_run_end(text, m.end())
        if end > start:
            s2, e2 = start, end
            for op, cl in (("(", ")"), ("[", "]"), ("{", "}")):
                k = s2 - 1
                while k >= 0 and text[k] == " ":
                    k -= 1
                t = e2
                while t < n and text[t] == " ":
                    t += 1
                if k >= 0 and text[k] == op and t < n and text[t] == cl:
                    s2, e2 = k, t + 1
                    break
            spans.append((s2, e2))
            i = max(e2, m.end())
    return spans


def _is_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|") and "|" in line[1:]


def _render_markdown(doc: Document, md: str):
    lines = md.splitlines()
    i = 0
    ordered_idx = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Inline math: convert one-line $...$ expressions to Word equations.
        if line.count("$") >= 2 and "$$" not in line:
            p = doc.add_paragraph()
            parts = re.split(r"\$([^$]+)\$", line)
            for idx, part in enumerate(parts):
                if not part:
                    continue
                if idx % 2:
                    _append_equation(p, part)
                else:
                    _add_runs(p, part)
            i += 1
            continue

        # Tabel markdown
        if _is_table_row(line) and i + 1 < len(lines) and re.match(
            r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]
        ):
            header_cells = [c.strip() for c in line.strip().strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and _is_table_row(lines[j]):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            table = doc.add_table(rows=1, cols=len(header_cells))
            table.style = "Table Grid"
            for col, cell in enumerate(header_cells):
                table.rows[0].cells[col].text = cell
                table.rows[0].cells[col].paragraphs[0].runs[0].bold = True
            for row_data in rows:
                cells = table.add_row().cells
                for col in range(len(header_cells)):
                    if col < len(row_data):
                        cells[col].text = row_data[col]
                    else:
                        cells[col].text = ""
            doc.add_paragraph()
            i = j
            continue

        # Equation Word (OMML): baris yang memuat $$...$$
        if "$$" in line:
            m = re.match(r"^([-*])\s+", line)
            style = "List Bullet" if m else None
            base = re.sub(r"^[-*]\s+", "", line)
            p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
            parts = base.split("$$")
            for idx, seg in enumerate(parts):
                if idx % 2 == 0:
                    if seg.strip():
                        _add_runs(p, seg)
                else:
                    _append_equation(p, seg)
            if re.fullmatch(r"\$\$.*\$\$", base.strip()):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ordered_idx = 0
            i += 1
            continue

        # Equation LaTeX telanjang (tanpa $): baris seperti
        # `YmidX;∼;Normal(\beta_0 + \beta_1 x_1,\; \sigma^2)`.
        if (
            re.match(r"^[-*>#|\d]", line.strip()) is None
            and _line_is_display_math(line)
        ):
            p = doc.add_paragraph()
            if len(line.strip()) < 60:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _append_equation(p, line.strip())
            i += 1
            continue

        # Heading
        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph()
            run = p.add_run(m.group(2).strip())
            run.bold = True
            if level == 2:
                run.font.size = Pt(14)
            elif level == 3:
                run.font.size = Pt(12)
            i += 1
            continue

        # Bullet list
        if re.match(r"^[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, re.sub(r"^[-*]\s+", "", line))
            i += 1
            ordered_idx = 0
            continue

        # Ordered list
        mo = re.match(r"^(\d+)[.)]\s+(.*)$", line)
        if mo:
            num = mo.group(1).lstrip("0") or "0"
            p = doc.add_paragraph()
            p.add_run(f"{num}. ")
            _add_runs(p, mo.group(2))
            ordered_idx = 1
            i += 1
            continue

        ordered_idx = 0
        # Empty line → skip; blockquote → plain paragraph
        if not line.strip():
            i += 1
            continue
        # Horizontal rule --- / *** → spacer paragraph
        if re.match(r"^[-*=_]{3,}\s*$", line.strip()):
            doc.add_paragraph()
            i += 1
            continue
        p = doc.add_paragraph()
        _add_runs(p, line.strip())
        i += 1


def _clean_soal(md: str) -> str:
    """Buang metadata lampiran/transkripsi dari soal, pertahankan isi transkripsi."""
    lines = md.splitlines()
    out: list[str] = []
    skip_attachments = False
    for line in lines:
        stripped = line.strip()
        if re.match(r"^##\s+Lampiran\s*:?", stripped):
            skip_attachments = True
            continue
        if re.match(r"^##\s+Isi\s+lampiran\b", stripped, re.IGNORECASE):
            continue
        if stripped == "# Transkrip Soal":
            continue
        if skip_attachments:
            if re.match(r"^[-*]\s+", stripped) or not stripped:
                continue
            skip_attachments = False
        out.append(line)
    return "\n".join(out)


def build_docx(
    *,
    jawaban_md: str,
    soal_text: str,
    meta: dict,
    out_docx: Path,
    include_soal: bool = True,
) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        q = qn(attr)
        if q in rf.attrib:
            del rf.attrib[q]
    rf.set(qn("w:eastAsia"), "Times New Roman")

    # Header identitas
    for label, key in (("Nama", "nama"), ("NIM", "nim"), ("Prodi", "prodi"), ("Matkul", "matkul")):
        p = doc.add_paragraph()
        r = p.add_run(f"{label} : ")
        r.bold = True
        p.add_run(str(meta.get(key, "")))

    # Judul
    title = f"{meta.get('kind_label', '')} {meta.get('display_index', '')} {meta.get('matkul', '')}".strip()
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = tp.add_run(title)
    tr.bold = True
    tr.font.size = Pt(13)

    # Soal
    if include_soal and soal_text.strip():
        doc.add_paragraph()
        sp = doc.add_paragraph()
        sr = sp.add_run("Soal")
        sr.bold = True
        sr.font.size = Pt(14)
        _render_markdown(doc, _clean_soal(soal_text))

    # Jawab (dari markdown opencode)
    try:
        jawab_idx = jawaban_md.lower().find("## jawab")
        body = jawaban_md[jawab_idx:] if jawab_idx != -1 else jawaban_md
    except Exception:  # noqa: BLE001
        body = jawaban_md
    _render_markdown(doc, body)

    _apply_hanging_indent_refs(doc)

    out_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_docx))
    return out_docx


def _apply_hanging_indent_refs(doc: Document):
    """Referensi (setelah 'Daftar Pustaka') diberi hanging indent ala APA."""
    from docx.shared import Cm

    found = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.lower().startswith("daftar pustaka"):
            found = True
            continue
        if found and txt:
            pf = p.paragraph_format
            pf.left_indent = Cm(0.63)
            pf.first_line_indent = Cm(-0.63)


def _convert_to_doc(docx_path: Path, doc_path: Path) -> bool:
    try:
        import win32com.client  # type: ignore

        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        try:
            wd = word.Documents.Open(str(docx_path.resolve()))
            wd.SaveAs2(str(doc_path.resolve()), FileFormat=0)  # 0 = Word 97-2003 .doc
            wd.Close(False)
        finally:
            word.Quit()
        return True
    except Exception as exc:  # noqa: BLE001
        print(f"  ! Konversi ke .doc dilewati ({exc}). File .docx dipertahankan.")
        return False


def save_doc(jawaban_md: str, soal_text: str, meta: dict, out_dir: Path) -> tuple[Path, Path]:
    """Simpan jawaban sebagai satu file .docx final (equation OMML asli).

    Equation asli (Word Equation) hanya dapat disimpan dalam format OOXML
    (.docx); konversi ke .doc lama akan mengubah equation menjadi gambar.
    Return (docx_path, docx_path) agar pemanggil tetap API tuple.
    """
    base = meta.get("file_base", "jawaban")
    docx_path = out_dir / f"{base}.docx"
    doc_path = out_dir / f"{base}.doc"
    doc_path.unlink(missing_ok=True)

    build_docx(
        jawaban_md=jawaban_md,
        soal_text=soal_text,
        meta=meta,
        out_docx=docx_path,
    )
    return docx_path, docx_path