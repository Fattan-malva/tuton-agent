"""Uji "Format Jawaban" (Form Soal), identitas dokumen, dan folder `output/`.

Semuanya tidak butuh Moodle dan tidak membebani model: `build_file_prompt`
adalah fungsi murni, dokumen .docx dirender lokal dengan python-docx, dan
pemeriksaan UI cuma membandingkan className di berkas .tsx.

Yang dijaga di sini:

  1. Field opsional benar-benar opsional. Kosong -> prompt persis seperti
     versi tanpa fitur ini, dan dokumen tetap memakai template standar.
  2. Berkas contoh format dipakai sebagai DOKUMEN DASAR, bukan ditiru. Kalau
     ini di-hardcode style-nya, hasil dokumen bisa terlihat mirip tapi
     geometrinya salah -- dan geometrilah yang paling kelihatan ke censor.
  3. Dua dokumen yang dirender tanpa berkas format tidak boleh saling
     menimpa, dan tidak boleh ikut berubah karena ada berkas format lain.
  4. Semester dan UT Daerah benar-benar sampai ke dokumen, dan boleh dikosongkan.
  5. `output/` dibuat otomatis dan seluruh isinya terabaikan git.
  6. Class `pixel-*` tidak ditimpa utility Tailwind yang justru tidak
     berlaku --itusalah penyebab ikon search menabrak teksnya.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, ".")

import config  # noqa: E402
from dotenv import dotenv_values  # noqa: E402
from generator.docx import save_doc  # noqa: E402
from generator.prompt import build_file_prompt  # noqa: E402

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

PASS: list[str] = []
FAIL: list[tuple[str, str]] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    if condition:
        PASS.append(name)
        print(f"  PASS  {name}")
    else:
        FAIL.append((name, detail))
        print(f"  FAIL  {name}  {detail}")


JAWABAN = """## Jawaban Mahasiswa

### Analisis

Basis data relasi menyimpan data dalam tabel.

## Daftar Pustaka

1. Connolly, T., & Begg, C. (2015). *Database Systems* (6th ed.). Pearson.
2. Elmasri, R., & Navathe, S. B. (2016). *Fundamentals of Database Systems*.
3. https://doi.org/10.1145/3292500.3330718
"""

META = {
    "nama": "Mahasiswa Uji",
    "nim": "2230100",
    "semester": "3",
    "ut_daerah": "Jakarta",
    "matkul": "Basis Data 64",
    "kind_label": "Tugas",
    "display_index": 1,
    "file_base": "uji",
}

# python-docx menyimpan ukuran halaman dalam twips, jadi EMU yang diberikan
# tidak kembali persis. 1000 EMU ~= 0,0028 cm: jauh di bawah pembeda yang
# berarti secara visual, tapi cukup untuk membedakan dokumen yang benar.
TOL = 1000


def dekat(a, b) -> bool:
    return a is not None and b is not None and abs(int(a) - int(b)) <= TOL


def _tulis_dotenv(path: Path, isi: str) -> bool:
    """Tulis berkas .env sementara. `False` kalau gagal (folder tidak ada)."""
    try:
        path.write_text(isi, encoding="utf-8")
        return True
    except OSError:
        return False


def _baca_satu_kunci(isi: str, kunci: str) -> str | None:
    """Baca satu variabel dari teks .env, persis seperti python-dotenv."""
    sementara = Path("output") / ".testenv_probe"
    if not _tulis_dotenv(sementara, isi):
        return None
    try:
        return dotenv_values(sementara).get(kunci)
    finally:
        sementara.unlink(missing_ok=True)


def build_prompt(**extra) -> str:
    return build_file_prompt(
        work_kind="tugas",
        index=1,
        course_name="Basis Data 64",
        section_num=1,
        activity_title="Buka tugas",
        soal_path=Path("output/soal.md"),
        attachment_dir=Path("output/lampiran"),
        lampiran=["tugas1.pdf"],
        jawaban_path=Path("output/jawaban.md"),
        **extra,
    )


def _buat_berkas_contoh(path: Path) -> None:
    """Berkas .docx dengan geometri dan font yang beda dari template standar.

    Font dipasang di tingkat RUN, bukan hanya di style `Normal` -- itu yang
    dilakukan template asli (style `Normal`-nya menulis Arial, tapi tiap run
    membawa override Times New Roman). `_template_body_font` sengaja mencari
    run, bukan mempercayai style, jadi fixture harus mengikuti pola itu juga.
    """
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = 30_000_000, 20_000_000
    section.left_margin = section.right_margin = 500_000
    doc.styles["Normal"].font.name = "Arial"  # sengaja beda, harus diabaikan
    run = doc.add_paragraph("Contoh jawaban dengan isi yang cukup panjang "
                            "supaya jadi run Majority.").runs[0]
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    doc.save(str(path))


def test_prompt_optional(scratch: Path) -> None:
    print("\n[1] Field Format Jawaban benar-benar opsional")
    polos = build_prompt()
    check("tanpa format tidak ada blok format", "Format jawaban (WAJIB" not in polos)
    check("tanpa format aturan Daftar Pustaka tetap ada", "Daftar Pustaka" in polos)

    catatan = build_prompt(format_note="Pakai judul ABSTRAK, tanpa tabel.")
    check("keterangan saja memunculkan blok format",
          "Format jawaban (WAJIB" in catatan)
    check("keterangan masuk prompt apa adanya", "Pakai judul ABSTRAK" in catatan)
    check("keterangan tidak dikira berkas contoh",
          "Berkas contoh format:" not in catatan)
    check("bagian lama tidak berubah saat menambah format",
          catatan.count("## Gaya jawaban") == polos.count("## Gaya jawaban"))

    berkas = build_prompt(format_path=scratch / "contoh_format.docx")
    check("berkas contoh disebut di prompt", "contoh_format.docx" in berkas)
    check("tidak ada kata keterangan tambahan saat kosong",
          "Keterangan tambahan" not in berkas)
    check("pembagian peran dijelaskan ke model",
          "PYTHON" not in berkas.upper().split("BERKAS CONTOH")[0][-400:]
          and "pipeline" in berkas.lower())

    keduanya = build_prompt(
        format_path=scratch / "contoh_format.docx",
        format_note="Pakai judul ABSTRAK.",
    )
    check("berkas dan keterangan bisa dipakai bersamaan",
          "contoh_format.docx" in keduanya and "Pakai judul ABSTRAK" in keduanya)
    check("konflik antara keduanya dijelaskan", "saling berbeda" in keduanya)


def test_docx_template(scratch: Path) -> None:
    print("\n[2] Berkas format dipakai sebagai dokumen dasar")
    contoh = scratch / "contoh_format.docx"
    _buat_berkas_contoh(contoh)

    # `file_base` wajib berbeda: `save_doc` menamai berkas dari `file_base`,
    # jadi panggilan kedua akan menimpa yang pertama kalau sama.
    _, std_path = save_doc(JAWABAN, "", {**META, "file_base": "std"}, scratch)
    _, custom_path = save_doc(
        JAWABAN, "", {**META, "file_base": "custom"}, scratch, template=contoh
    )
    check("dua berkas terpisah", std_path != custom_path and std_path.is_file()
          and custom_path.is_file())

    std, custom = Document(std_path), Document(custom_path)
    gs, gc = std.sections[0], custom.sections[0]
    check("ukuran halaman ikut berkas contoh",
          dekat(gc.page_width, 30_000_000) and dekat(gc.page_height, 20_000_000),
          f"{(gc.page_width, gc.page_height)}")
    check("margin ikut berkas contoh", dekat(gc.left_margin, 500_000),
          f"{gc.left_margin}")
    check("dokumen tanpa format tidak ikut berubah",
          not (dekat(gs.page_width, 30_000_000) and dekat(gs.page_height, 20_000_000)),
          f"{(gs.page_width, gs.page_height)}")

    def font_body(doc) -> str | None:
        """Font yang benar-benar dipakai isi dokumen, bukan yang tertulis di style.

        Diuji lewat `_apply_base_typography`: isi dokumen hasil tidak lagi
        membawa override sendiri, jadi font efektifnya ada di style `Normal`.
        """
        return doc.styles["Normal"].font.name

    check("font body ikut berkas contoh", font_body(custom) == "Courier New",
          f"font={font_body(custom)}")
    check("style Normal yang menyesat di template diabaikan",
          font_body(custom) != "Arial", f"font={font_body(custom)}")

    print("\n[3] Batas keras daftar pustaka")
    # Bagian 4-5 ada di test_identitas_dokumen, yang dijalankan sesudah bagian 3.
    check("bawaan TUTON_MAX_PUSTAKA = 5", config.Config.TUTON_MAX_PUSTAKA == 5,
          str(config.Config.TUTON_MAX_PUSTAKA))
    ref = [p.text for p in std.paragraphs if p.text.strip()[:1].isdigit()]
    check("daftar pustaka di docx tidak melebihi 5", len(ref) <= 5, f"{len(ref)}")


def test_output_dir() -> None:
    print("\n[7] Folder output/ dibuat otomatis")
    root = config.OUTPUT_DIR
    backup = root.with_name(f"{root.name}__backup_uji")
    if backup.exists():
        shutil.rmtree(backup)
    ada_awal = root.exists()
    if ada_awal:
        shutil.move(str(root), str(backup))

    try:
        check("output/ tidak ada sebelum dipanggil", not root.exists())
        config.ensure_output_dirs()
        check("output/ dibuat", root.is_dir())
        check("output/.jobs dibuat", config.OUTPUT_JOBS_DIR.is_dir())
        check("output/.cache dibuat", config.OUTPUT_CACHE_DIR.is_dir())
        check("subfolder benar-benar di dalam output/",
              config.OUTPUT_JOBS_DIR.parent == root
              and config.OUTPUT_CACHE_DIR.parent == root)

        sentinel = root / "sentinel.txt"
        sentinel.write_text("x", encoding="utf-8")
        config.ensure_output_dirs()
        check("panggilan kedua tidak merusak isi (idempoten)", sentinel.is_file())

        nested = root / "MataKuliah_Uji" / "sesi1"
        nested.mkdir(parents=True, exist_ok=True)
        (nested / "jawaban_tugas_1.docx").write_bytes(b"PK\x03\x04")
        check("bisa langsung menulis di folder yang baru dibuat",
              (nested / "jawaban_tugas_1.docx").is_file())
    finally:
        shutil.rmtree(root, ignore_errors=True)
        if ada_awal:
            shutil.move(str(backup), str(root))
    check("folder keluaran asli dipulihkan", root.is_dir())


def test_gitignore() -> None:
    print("\n[8] Semua isi output/ diabaikan git")

    def diabaikan(path: str) -> bool:
        r = subprocess.run(["git", "check-ignore", "-q", "--", path], capture_output=True)
        return r.returncode == 0

    # Satu per satu. Satu pola yang terlewat berarti file itu bisa ikut ter-push
    # tanpa terlihat, dan isinya memuat Nama/NIM.
    contoh = [
        "output/soal.md",
        "output/state.json",
        "output/.jobs/solve_1.docx",
        "output/.cache/pages/x.html",
        "output/.cache/md/x.md",
        "output/.cache/transkrip/abc.txt",
        "output/MK_64/_petak/sesi1.md",
        "output/MK_64/sesi1/jawaban_tugas_1.md",
        "output/MK_64/sesi1/referensi_tugas_1.md",
        "output/MK_64/sesi1/MK_Tugas1.docx",
        "output/MK_64/sesi1/_format/contoh.docx",
        "output/MK_64/sesi1/lampiran/tugas1.pdf",
    ]
    for path in contoh:
        check(f"diabaikan: {path}", diabaikan(path))

    tracked = [
        l for l in subprocess.run(
            ["git", "ls-files", "output"], capture_output=True, text=True,
            encoding="utf-8", errors="replace",
        ).stdout.splitlines() if l.strip()
    ]
    check("tidak ada file output/ yang sudah ter-track", not tracked, str(tracked[:3]))


def test_identitas_dokumen() -> None:
    print("\n[4] Semester dan UT Daerah sampai ke dokumen")
    meta = {
        "nama": "Mahasiswa Uji",
        "nim": "2230100",
        "semester": "Semester 2",
        "ut_daerah": "Jakarta",
        "matkul": "Basis Data 64",
        "kind_label": "Tugas",
        "display_index": 1,
        "file_base": "ident",
    }
    out = Path("output") / ".testformat"
    out.mkdir(parents=True, exist_ok=True)
    _, path = save_doc(JAWABAN, "", meta, out)
    doc = Document(path)
    isi_tabel = {
        row.cells[0].text.strip(): row.cells[1].text.strip()
        for row in doc.tables[0].rows
        if len(row.cells) >= 2
    }
    check("Semester masuk tabel identitas",
          isi_tabel.get("Semester") == "Semester 2", str(isi_tabel.get("Semester")))
    check("UT Daerah masuk tabel identitas",
          isi_tabel.get("UT Daerah") == "Jakarta", str(isi_tabel.get("UT Daerah")))

    # Baris kosong harus DILEWATI, bukan dicetak sebagai label tanpa nilai:
    # "Semester |" yang menggantung terlihat seperti tabel gagal terisi.
    meta_kosong = {**meta, "semester": "", "ut_daerah": "", "file_base": "kosong"}
    _, path_kosong = save_doc(JAWABAN, "", meta_kosong, out)
    isi_kosong = {
        row.cells[0].text.strip()
        for row in Document(path_kosong).tables[0].rows
    }
    check("Semester kosong tidak jadi baris menggantung", "Semester" not in isi_kosong,
          str(sorted(isi_kosong)))
    check("UT Daerah kosong tidak jadi baris menggantung", "UT Daerah" not in isi_kosong,
          str(sorted(isi_kosong)))

    print("\n[5] Semester dan UT Daerah bisa disimpan dari Settings")
    server = Path("server.py").read_text(encoding="utf-8")
    check("/api/config mengirim semester", '"semester": Config.SEMESTER' in server)
    check("/api/config mengirim ut_daerah", '"ut_daerah": Config.UT_DAERAH' in server)
    check("POST /api/config menulis SEMESTER", '"SEMESTER":' in server)
    check("POST /api/config menulis UT_DAERAH", '"UT_DAERAH":' in server)
    settings = Path("web/src/components/Settings.tsx").read_text(encoding="utf-8")
    check("form punya field semester", 'name="semester"' in settings)
    check("form punya field ut_daerah", 'name="ut_daerah"' in settings)

    print("\n[6] Nama variabel di .env.example bisa diurai")
    # `UT Daerah=Jakarta` SEHARUSNYA gagal: nama variabel dengan spasi tidak
    # bisa diurai python-dotenv, barisnya dilewati diam-diam, dan isinya
    # tidak pernah sampai ke dokumen -- tanpa error yang terlihat.
    check("nama variabel ber-spasi memang tidak terbaca (dasar masalahnya)",
          _baca_satu_kunci("UT Daerah=Jakarta\n", "UT Daerah") is None)
    check("nama variabel tanpa spasi terbaca normal",
          _baca_satu_kunci("UT_DAERAH=Jakarta\n", "UT_DAERAH") == "Jakarta")

    isi_contoh = Path(".env.example").read_text(encoding="utf-8")
    check(".env.example memakai SEMESTER (bukan Semester)",
          re.search(r"^SEMESTER=", isi_contoh, re.MULTILINE) is not None
          and re.search(r"^Semester=", isi_contoh, re.MULTILINE) is None)
    check(".env.example memakai UT_DAERAH (bukan 'UT Daerah')",
          re.search(r"^UT_DAERAH=", isi_contoh, re.MULTILINE) is not None
          and re.search(r"^UT Daerah=", isi_contoh, re.MULTILINE) is None)

    # Setiap kunci di .env.example harus bisa diurai dan dibaca config.py.
    kunci_bermasalah = []
    for baris in isi_contoh.splitlines():
        if "=" not in baris or baris.lstrip().startswith("#"):
            continue
        kunci = baris.split("=", 1)[0].strip()
        if not kunci or " " in kunci:
            kunci_bermasalah.append(baris)
    check("tidak ada kunci ber-spasi di .env.example", not kunci_bermasalah,
          str(kunci_bermasalah[:3]))

    # config.py harus tetap membaca ejaan lama, supaya .env milik pengguna
    # yang terlanjur salah edit tidak harus diperbaiki manual.
    cfg = Path("config.py").read_text(encoding="utf-8")
    check("config.py menerima ejaan lama lewat _env_any",
          "_env_any(\"SEMESTER\"" in cfg and "_env_any(\"UT_DAERAH\"" in cfg)
    check("Settings membuang kunci ejaan lama saat menyimpan",
          "_ENV_LEGACY_KEYS" in server)


def test_kelas_pixel() -> None:
    """Pagar untuk bug ikon search menabrak teks.

    Akar masalahnya: aturan `.pixel-*` di `globals.css` ditulis DI LUAR
    `@layer`. Di kaskade CSS, aturan tanpa layer selalu menang atas utility
    Tailwind yang berada di `@layer utilities`. Jadi `className="pixel-input
    pl-8"` terlihat seperti memberi padding kiri 2rem, padahal `padding: 0
    12px` dari `.pixel-input` tetap yang berlaku -- teks mulai di 12px dan
    menabrak ikon yang duduk di 10px.

    Pemeriksaan ini tidak menangkap hasil visual, tapi menangkap NIAT: siapa
    pun yang menulis utility padding/min-height di samping class `pixel-input`
    akan gagal di sini, bukan diam-diam salah di layar.
    """
    print("\n[9] Class pixel-* tidak ditimpa utility yang tidak berlaku")
    pola_kelas = re.compile(r"""className=(?:\{)?\s*["'`]([^"'`]+)["'`]""")
    # Utility yang kalah telak terhadap `.pixel-input` (padding + min-height).
    utility_berbahaya = re.compile(
        r"(?:^|\s)(?:p|px|py|pt|pb|pl|pr|min-h|max-h|h)(?:-\d|\[)"
    )

    offender: dict[str, str] = {}
    for tsx in sorted(Path("web/src").rglob("*.tsx")):
        teks = tsx.read_text(encoding="utf-8")
        for cocok in pola_kelas.finditer(teks):
            kelas = cocok.group(1)
            if "pixel-input" not in kelas:
                continue
            # `.pixel-textarea` sudah mengurus tinggi dan padding vertikalnya.
            if "pixel-textarea" in kelas or "pixel-input-icon" in kelas:
                continue
            ketemu = utility_berbahaya.search(kelas)
            if ketemu:
                offender[f"{tsx.name}:{teks[:cocok.start()].count(chr(10)) + 1}"] = (
                    f"{kelas.strip()} (utility {ketemu.group(0).strip()})"
                )

    check("tidak ada pixel-input + utility padding/min-height yang tak berlaku",
          not offender, "; ".join(f"{k} -> {v}" for k, v in offender.items()))

    # Ikon di dalam input: harus ada padding kiri yang benar-benar berlaku.
    # Dicari di dalam `className`, bukan di seluruh berkas -- nama `pl-8` juga
    # muncul di komentar penjelas, dan komentar memang harus menyebutnya.
    settings_kelas = [
        m.group(1)
        for m in pola_kelas.finditer(
            Path("web/src/components/Settings.tsx").read_text(encoding="utf-8")
        )
    ]
    check("input model pakai pixel-input-icon, bukan utility pl-*",
          any("pixel-input pixel-input-icon" in k for k in settings_kelas)
          and not any("pl-" in k for k in settings_kelas),
          str([k for k in settings_kelas if "pixel-input" in k]))

    css = Path("web/src/app/globals.css").read_text(encoding="utf-8")
    check("kelas pixel-input-icon benar-benar ada di CSS", ".pixel-input-icon" in css)
    check("kelas pixel-textarea benar-benar ada di CSS", ".pixel-textarea" in css)


def test_chrome_mobile_tidak_tertimbun() -> None:
    """Header/nav mobile harus punya z-index di atas `.pixel-world`.

    `.pixel-world` (page.tsx) adalah `position: fixed; z-index: 0` dengan
    background OPAQUE. Dalam urutan paint CSS, elemen *positioned* selalu
    menggambar di atas elemen in-flow yang statis -- jadi chrome mobile yang
    tidak punya positioning/z-index akan terkubur sepenuhnya di baliknya:
    tidak terlihat, dan tidak bisa diklik (`pixel-world` juga
    `pointer-events: none`).

    Bug ini pernah nyata: header mobile punya profil + logout tapi sama sekali
    tidak muncul karena `<aside>` (z-30) dan nav bawah (z-30) sudah aman,
    sedangkan header satu-satunya yang lupa. Karena itu diuji di sini, bukan
    hanya mengandalkan descubrimiento manual.
    """
    nav = Path("web/src/components/Navigation.tsx").read_text(encoding="utf-8")

    header = re.search(r"<header\b(.*?)>", nav, re.S)
    check("header mobile ditemukan di Navigation.tsx", header is not None)
    if header is None:
        return
    kelas = header.group(1)
    check("header mobile punya z-30 (di atas .pixel-world z-0)", "z-30" in kelas, kelas[:90])
    check("header mobile positioned (fixed/sticky/relative)",
          any(p in kelas for p in ("fixed", "sticky", "relative")), kelas[:90])
    check("header mobile disembunyikan di desktop (md:hidden)", "md:hidden" in kelas)

    aside = re.search(r"<aside\b(.*?)>", nav, re.S)
    check("sidebar desktop tetap z-30", aside is not None and "z-30" in aside.group(1))
    bottom = re.search(r"<nav\b[^>]*aria-label=\"Navigasi mobile\"", nav)
    check("nav bawah mobile tetap ada", bottom is not None)
    if bottom is not None:
        check("nav bawah mobile tetap z-30", "z-30" in bottom.group(0))

    # Padding konten wajib mengikuti, karena header-nya `fixed`.
    page = Path("web/src/app/page.tsx").read_text(encoding="utf-8")
    check("konten main memakai app-main-top-pad", "app-main-top-pad" in page)
    css = Path("web/src/app/globals.css").read_text(encoding="utf-8")
    check("variabel --app-header-h ada di CSS", "--app-header-h" in css)
    check("padding mobile dibatasi max-width (tidak mengalahkan md:pt-8)",
          "app-main-top-pad" in css and "max-width: 47.999rem" in css)
    check("tinggi header aman untuk notch iOS",
          "env(safe-area-inset-top" in css or "env(safe-area-inset-top" in nav)

    # Komentar CSS yang tidak tertutup membuat SEMUA aturan setelahnya ikut
    # ter-comment dan hilang diam-diam dari hasil build. Gejalanya sangat
    # menyesatkan: pemeriksaan sebelumnya hanya mencari kata kunci
    # `--app-header-h` di file sumber -- dan itu LOLOS padahal aturannya sudah
    # tidak pernah sampai ke CSS ter-build. Yang diuji di sini keseimbangan
    # pembuka/penutup, plus isi `:root` harus benar-benar ada di luar komentar.
    check("jumlah /* sama dengan */ di globals.css (tidak ada komentar tak tertutup)",
          css.count("/*") == css.count("*/"),
          f"/*={css.count('/*')} */={css.count('*/')}")
    tanpa_komentar = re.sub(r"/\*.*?\*/", "", css, flags=re.S)
    check("definisi --app-header-h berada di luar komentar",
          "--app-header-h:" in tanpa_komentar)
    check("aturan .app-main-top-pad berada di luar komentar",
          ".app-main-top-pad" in tanpa_komentar)


def main() -> int:
    scratch = Path("output") / ".testformat"
    if scratch.exists():
        shutil.rmtree(scratch)
    # Dibuat lewat helper resmi, bukan `mkdir` manual: kalau helper-nya rusak,
    # test ini ikut gagal dan itu memang yang ingin diuji.
    config.ensure_output_dirs()
    scratch.mkdir(parents=True, exist_ok=True)

    try:
        test_prompt_optional(scratch)
        test_docx_template(scratch)
        test_identitas_dokumen()
    finally:
        shutil.rmtree(scratch, ignore_errors=True)

    test_output_dir()
    test_gitignore()
    test_kelas_pixel()
    test_chrome_mobile_tidak_tertimbun()

    print(f"\nLULUS: {len(PASS)}   GAGAL: {len(FAIL)}")
    for name, detail in FAIL:
        print(f"  - {name}  {detail}")
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(main())
