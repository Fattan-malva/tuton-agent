from __future__ import annotations

import json
from pathlib import Path
import subprocess
import time
import zipfile

from config import Config, PROJECT_ROOT
from generator.opencode_runner import run_opencode
from moodle.ocr import is_image, ocr_image

_PDF_EXT = {".pdf"}
_XLSX_EXT = {".xlsx", ".xlsm", ".xls"}
_DOCX_EXT = {".docx"}

_TRANSCRIBE_PROMPT = """Kamu adalah agen transkripsi soal. Sebuah {TYPE} soal matematika terlampir pada pesan ini dan kamu bisa melihatnya langsung (modelmu mendukung vision).

TUGAS:
1. Baca seluruh isi {TYPE} yang terlampir.
2. Salin PERSIS seluruh teks soal ke file:
   {OUT}

ATURAN:
- Jangan menjawab soal, jangan berkomentar, jangan menganalisis.
- Salin hanya isi soal (judul, nomor soal, pertanyaan, instruksi).
- Tulis notasi matematika dengan jelas dalam teks polos/markdown:
  matriks -> [[a, b], [c, d]]; pangkat -> x^2; pecahan -> a/b; akar -> sqrt(...).
- Bagian yang tidak terbaca -> tulis [tidak terbaca].
- Tulis file hasil dalam satu kali operasi (buat file baru / timpa isinya).
- Di akhir, output terminal hanya satu baris: SELESAI"""


def _image_candidates() -> list[tuple[str, str]]:
    """Model vision untuk gambar: pilih otomatis yang mendukung input image."""
    return list(_candidates(pdf=False))


def _pdf_candidates() -> list[tuple[str, str]]:
    """Model vision untuk PDF: pilih otomatis yang mendukung input image+pdf."""
    return list(_candidates(pdf=True))


# ============================================================================
# Auto-deteksi model vision lewat `opencode models --verbose`.
# Model tidak di-hardcode karena daftar bisa berubah sewaktu-waktu (ada
# model yang hilang/tambah). Diprioritaskan: (1) prefer list dari config,
# (2) urutan yang dikembalikan opencode.
# ============================================================================
_models_cache: tuple[float, list[dict]] | None = None
_MODELS_CACHE_TTL = 120  # detik: hindari panggil CLI tiap attachment


def _parse_verbose_models(data: str) -> list[dict]:
    """Parse `opencode models --verbose`: setiap model = `id\n{json}\n`."""
    models: list[dict] = []
    pending_name: str | None = None
    depth = 0
    buf: list[str] = []
    for raw in data.splitlines():
        if pending_name is None:
            line = raw.strip()
            if not line:
                continue
            pending_name = line
            depth = 0
            buf = []
            continue
        if not buf and "{" not in raw:
            continue
        buf.append(raw)
        depth += raw.count("{") - raw.count("}")
        if depth == 0:
            try:
                meta = json.loads("\n".join(buf))
            except json.JSONDecodeError:
                meta = {}
            if meta:
                full = f"{meta.get('providerID', '')}/{meta.get('id', '')}".strip("/")
                meta["full_id"] = full
                models.append(meta)
            pending_name = None
    return models


def _fetch_models() -> list[dict]:
    global _models_cache
    now = time.time()
    if _models_cache is not None and now - _models_cache[0] < _MODELS_CACHE_TTL:
        return _models_cache[1]
    models: list[dict] = []
    try:
        from generator.opencode_runner import _resolve_opencode

        cmd = _resolve_opencode() + ["models", "--verbose"]
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
            cwd=str(PROJECT_ROOT),
        )
        data = (result.stdout or "") + (result.stderr or "")
        models = _parse_verbose_models(data)
    except Exception as exc:  # noqa: BLE001 - apapun blokirannya, lanjut fallback
        print(f"    ! gagal membaca daftar model opencode: {exc}")
    _models_cache = (now, models)
    return models


def _filter_vision_models(models: list[dict], *, pdf: bool) -> list[dict]:
    out: list[dict] = []
    for m in models:
        status = m.get("status") or "active"
        if status != "active":
            continue
        caps = m.get("capabilities") or {}
        inp = caps.get("input") or {}
        if not inp.get("image"):
            continue
        if pdf and not inp.get("pdf"):
            continue
        out.append(m)
    return out


def _match_model_id(pref: str, ids: list[str]) -> str | None:
    pref = pref.strip().strip("/")
    for mid in ids:
        if mid == pref or mid.endswith("/" + pref):
            return mid
    return None


def _ordered_vision_ids(*, pdf: bool) -> list[str]:
    available = _filter_vision_models(_fetch_models(), pdf=pdf)
    ids: list[str] = []
    for m in available:
        fid = m.get("full_id")
        if fid and fid not in ids:
            ids.append(fid)
    prefer = Config.OPENCODE_VISION_PREFER
    if prefer:
        ranked: list[str] = []
        for p in prefer:
            hit = _match_model_id(p, ids)
            if hit and hit not in ranked:
                ranked.append(hit)
        ids = ranked + [i for i in ids if i not in ranked]
    return ids


def _candidates(*, pdf: bool) -> list[tuple[str, str]]:
    """(model_id, variant) untuk transcriber. Variant hanya dipakai bila model
    menyediakan variant tsb (mis. `low` untuk reasoning effort)."""
    models = _filter_vision_models(_fetch_models(), pdf=pdf)
    variants_map = {
        m.get("full_id", ""): set(m.get("variants") or {}) for m in models
    }
    variant = Config.OPENCODE_VISION_VARIANT
    cands: list[tuple[str, str]] = []
    for mid in _ordered_vision_ids(pdf=pdf):
        v = variant if variant and variant in variants_map.get(mid, set()) else ""
        cands.append((mid, v))
    return cands


def _transcribe_with_model(
    path: Path, out_file: Path, model: str, variant: str
) -> str:
    """Jalankan agent transcriber (model vision) untuk menyalin soal ke file."""
    out_file.unlink(missing_ok=True)
    kind = "PDF" if path.suffix.lower() in _PDF_EXT else "gambar"
    prompt = (
        _TRANSCRIBE_PROMPT.replace("{TYPE}", kind)
        .replace("{OUT}", str(out_file))
    )
    try:
        result = run_opencode(
            prompt,
            agent="transcriber",
            attach=[str(path)],
            model=model,
            variant=variant,
            timeout=Config.TUTON_TIMEOUT_TRANSCRIBE,
        )
    except TimeoutError:
        print(f"    ! transkripsi {path.name} timeout di model {model}")
        out_file.unlink(missing_ok=True)
        return ""
    if result.returncode != 0:
        print(f"    ! transkripsi {path.name} gagal di model {model} (exit {result.returncode})")
    if out_file.exists() and out_file.stat().st_size > 20:
        text = out_file.read_text(encoding="utf-8").strip()
        if text:
            return text
    out_file.unlink(missing_ok=True)
    return ""


def _transcribe_image(path: Path, out_file: Path) -> str:
    """Vision-model transcription, lalu fallback easyocr."""
    for model, variant in _image_candidates():
        text = _transcribe_with_model(path, out_file, model, variant)
        if text and len(text) >= 10:
            return text
    text = ocr_image(path)
    if text:
        return "(fallback easyocr)\n" + text
    return ""


def _transcribe_pdf(path: Path, out_file: Path) -> str:
    """Vision-model transcription, lalu fallback pymupdf."""
    for model, variant in _pdf_candidates():
        text = _transcribe_with_model(path, out_file, model, variant)
        if text and len(text) >= 10:
            return text
    try:
        import fitz

        doc = fitz.open(str(path))
        parts = [pg.get_text().strip() for pg in doc]
        text = "\n".join(p for p in parts if p).strip()
        if text:
            return "(fallback pymupdf)\n" + text
        rendered: list[str] = []
        for index, page in enumerate(doc, 1):
            image_path = path.parent / f".page_{path.stem}_{index}.png"
            page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).save(str(image_path))
            try:
                page_text = _transcribe_image(image_path, out_file)
                if page_text:
                    rendered.append(page_text)
            finally:
                image_path.unlink(missing_ok=True)
        if rendered:
            return "\n\n".join(rendered)
    except Exception as exc:  # noqa: BLE001
        print(f"    ! pymupdf gagal untuk {path.name}: {exc}")
    return ""


def _extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True, read_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        head = f"[Sheet: {ws.title}]"
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = [
                str(c).strip() for c in row if c is not None and str(c).strip()
            ]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(head + "\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(blocks)


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    blocks: list[str] = []
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    if paras:
        blocks.append("\n".join(paras))
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            blocks.append("\n".join(rows))
    return "\n\n".join(blocks)


def _extract_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace").strip()


def _extract_docx_images(path: Path, work_dir: Path) -> list[str]:
    """Transcribe images embedded inside a DOCX attachment."""
    results: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/media/"):
                continue
            image_path = work_dir / f".embedded_{Path(name).name}"
            image_path.write_bytes(archive.read(name))
            try:
                text = _transcribe_image(image_path, work_dir / f"transkrip_{image_path.stem}.md")
                if text:
                    results.append(text)
            finally:
                image_path.unlink(missing_ok=True)
    return results


def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in _XLSX_EXT:
        return _extract_xlsx(path)
    if ext in _DOCX_EXT:
        return _extract_docx(path)
    if ext in {".txt", ".csv", ".tsv", ".md", ".rtf"}:
        return _extract_text_file(path)
    return ""


def process_attachment(path: Path, work_dir: Path) -> str:
    """Ekstrak isi satu lampiran jadi teks (untuk soal.md).

    - Gambar/PDF: salin via agent transcriber (model vision), lalu fallback
      easyocr / pymupdf.
    - Excel/docx: ekstrak lewat Python.
    Kembalikan teks; string kosong bila tidak bisa (perlu cek manual).
    """
    ext = path.suffix.lower()
    if is_image(path) or ext in _PDF_EXT:
        if not path.is_file():
            return ""
        out_file = work_dir / f"transkrip_{path.stem}.md"
        if ext in _PDF_EXT:
            return _transcribe_pdf(path, out_file)
        return _transcribe_image(path, out_file)
    try:
        text = _extract_text(path).strip()
        if ext in _DOCX_EXT:
            embedded = _extract_docx_images(path, work_dir)
            if embedded:
                text = "\n\n".join(part for part in (text, *embedded) if part)
        return text
    except Exception as exc:  # noqa: BLE001
        print(f"    ! ekstraksi gagal untuk {path.name}: {exc}")
        return ""